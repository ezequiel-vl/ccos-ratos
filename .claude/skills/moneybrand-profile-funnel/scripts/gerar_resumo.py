"""Gera o resumo da Fase 1 do aluno em Word (.docx), pronto pra abrir no Google Docs.

Uso:
    python scripts/gerar_resumo.py dados.json resumo-fase-1.docx

Formato do JSON:
{
  "titulo": "Resumo da Fase 1 · Moneybrand + Oferta",
  "aluno": "Nome do aluno (opcional)",
  "partes": [
    {"titulo": "Uma Coisa Só", "campos": [["Problema", "..."], ["Pessoa", "..."], ["Bandeira", "..."]]},
    {"titulo": "Solução Nova, Problema Antigo", "campos": [["Mecanismo", "..."]]},
    {"titulo": "Carta Convite", "campos": [["Headline", "..."]]},
    {"titulo": "Validação", "campos": [["Mensagens enviadas", "..."]]}
  ],
  "carta_convite_texto": "Opcional. Texto completo da carta, com parágrafos separados por linha em branco."
}

Precisa da biblioteca python-docx. Se ela não estiver instalada, o script gera um .md com o mesmo conteúdo.
"""
import json
import sys
from pathlib import Path


def gerar_md(dados, saida):
    linhas = [f"# {dados.get('titulo', 'Resumo da Fase 1')}", ""]
    if dados.get("aluno"):
        linhas += [dados["aluno"], ""]
    for parte in dados.get("partes", []):
        linhas += [f"## {parte['titulo']}", ""]
        for rotulo, valor in parte.get("campos", []):
            linhas += [f"**{rotulo}:** {valor}", ""]
    if dados.get("carta_convite_texto"):
        linhas += ["## A carta convite", "", dados["carta_convite_texto"], ""]
    saida = Path(saida).with_suffix(".md")
    saida.write_text("\n".join(linhas), encoding="utf-8")
    print(f"python-docx não encontrado. Resumo gerado em Markdown: {saida}")


def gerar_docx(dados, saida):
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt, RGBColor

    doc = Document()
    estilo = doc.styles["Normal"]
    estilo.font.name = "Arial"
    estilo.font.size = Pt(11)
    for nome in ("Heading 1", "Heading 2"):
        s = doc.styles[nome]
        s.font.name = "Arial"
        s.font.bold = True
        s.font.color.rgb = RGBColor(0, 0, 0)

    cabecalho = doc.sections[0].header.paragraphs[0]
    cabecalho.text = "Moneybrand + Profile Funnel"
    cabecalho.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    cabecalho.runs[0].font.size = Pt(9)

    doc.add_heading(dados.get("titulo", "Resumo da Fase 1"), level=1)
    if dados.get("aluno"):
        doc.add_paragraph(dados["aluno"])

    def sombrear(celula, cor="F3F3F3"):
        tc = celula._tc.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), cor)
        tc.append(shd)

    for parte in dados.get("partes", []):
        doc.add_heading(parte["titulo"], level=2)
        campos = parte.get("campos", [])
        if not campos:
            continue
        tabela = doc.add_table(rows=0, cols=2)
        tabela.style = "Table Grid"
        for rotulo, valor in campos:
            linha = tabela.add_row().cells
            linha[0].text = str(rotulo)
            linha[0].paragraphs[0].runs[0].font.bold = True
            sombrear(linha[0])
            linha[1].text = str(valor)
        doc.add_paragraph()

    if dados.get("carta_convite_texto"):
        doc.add_heading("A carta convite", level=2)
        for bloco in str(dados["carta_convite_texto"]).split("\n\n"):
            if bloco.strip():
                doc.add_paragraph(bloco.strip())

    doc.save(saida)
    print(f"Resumo gerado: {saida}")


def main():
    if len(sys.argv) != 3:
        print(__doc__)
        sys.exit(1)
    dados = json.loads(Path(sys.argv[1]).read_text(encoding="utf-8"))
    try:
        import docx  # noqa: F401
    except ImportError:
        gerar_md(dados, sys.argv[2])
        return
    gerar_docx(dados, sys.argv[2])


if __name__ == "__main__":
    main()
